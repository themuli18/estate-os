"""Rauchtest für die deutschen Bank-Parser auf synthetischen Beispieldateien (keine echten Daten).
Aufruf: python3 tests/test_bankparsers_de.py   (benötigt openpyxl; pdftotext nur für echte PDFs)"""
import shutil, subprocess, sys, tempfile
from decimal import Decimal
from pathlib import Path

KIT = Path(__file__).resolve().parent.parent
TOOLS = KIT / "template-de" / "Konto" / "_tools"
sys.path.insert(0, str(TOOLS))
import camt053, csv_de, commerzbank_pdf as cb  # noqa: E402
from transaktion import betrag_de, sepa_felder  # noqa: E402

F = KIT / "tests" / "fixtures-de"
assert betrag_de("1.234,56-") == Decimal("-1234.56") and betrag_de("+12,30") == Decimal("12.30")
assert sepa_felder("EREF+X1 MREF+M-2 CRED+DE00ZZZ1 SVWZ+Text hier")["mandat"] == "M-2"

assert csv_de.erkenne(F / "Sparkasse/sparkasse_beispiel.csv") == "sparkasse"
assert csv_de.erkenne(F / "DKB/dkb_beispiel.csv") == "dkb"
assert csv_de.erkenne(F / "ING/ing_beispiel.csv") == "ing_de"
assert csv_de.erkenne(F / "CAMT/camt_beispiel.xml") == "camt053"

sp = csv_de.sparkasse(F / "Sparkasse/sparkasse_beispiel.csv")
assert sp[0].betrag == Decimal("-23.40") and sp[0].mandatsreferenz == "VS-12345678" and sp[0].referenz == ""
assert sp[2].betrag == Decimal("1250.00")
dkb = csv_de.dkb(F / "DKB/dkb_beispiel.csv")
assert dkb[0].konto == "DE00120300000000000002" and dkb[0].gegenpartei == "Beispiel Stadtwerke GmbH"
assert dkb[0].glaeubiger_id == "DE00ZZZ00000000003"
ing = csv_de.ing_de(F / "ING/ing_beispiel.csv")
assert ing[0].konto == "DE00500105170000000003" and ing[0].mandatsreferenz == "ABO-9"
assert ing[0].glaeubiger_id == "DE00ZZZ00000000004" and ing[0].verwendungszweck == "Abo Februar"
camt, a = camt053.parse(F / "CAMT/camt_beispiel.xml")
assert camt[0].betrag == Decimal("-95.00") and a[0]["stimmt"] is True

giro, p = cb.girokonto((F / "commerzbank_giro_beispiel.txt").read_text(), "giro.pdf")
assert len(giro) == 2 and giro[0].betrag == Decimal("-23.40") and giro[1].betrag == Decimal("1250.00")
assert giro[0].mandatsreferenz == "VS-12345678" and giro[0].glaeubiger_id == "DE00ZZZ00000000001"
assert giro[0].gegenkonto == "DE00200000000000000099" and giro[0].referenz == ""
assert p["stimmt"] is True and p["konto"] == "DE00600000000000000006"
karte, p = cb.kreditkarte((F / "commerzbank_karte_beispiel.txt").read_text(), "karte.pdf")
assert [t.betrag for t in karte] == [Decimal("-45.10"), Decimal("-23.40"), Decimal("100.00")]
assert karte[1].extra["waehrung"] == "GBP" and karte[0].konto == "Mastercard …1234"
assert p["stimmt"] is True and p["id"] == "2026-01-24"

with tempfile.TemporaryDirectory() as d:
    konto = Path(d) / "Konto"
    shutil.copytree(TOOLS, konto / "_tools")
    for sub in ("Sparkasse", "DKB", "ING", "CAMT"):
        shutil.copytree(F / sub, konto / sub)
    out = subprocess.run([sys.executable, str(konto / "_tools/baue_uebersicht.py")], capture_output=True, text=True)
    print(out.stdout, out.stderr)
    assert out.returncode == 0
    from openpyxl import load_workbook
    wb = load_workbook(konto / "Transaktionen.xlsx")
    assert wb.sheetnames == ["Buchungen", "Wiederkehrend", "Je Monat", "Prüfung"]
    wied = [r[0] for r in wb["Wiederkehrend"].iter_rows(min_row=2, values_only=True)]
    assert "Beispiel Stadtwerke GmbH" in wied and "Beispiel Versicherung AG" in wied, wied
    kat = {r[3]: r[4] for r in wb["Buchungen"].iter_rows(min_row=2, values_only=True)}
    assert kat["Eigenes Konto"].startswith("Intern")
print("OK – alle deutschen Bank-Parser-Tests bestanden")
