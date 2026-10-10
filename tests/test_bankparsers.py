"""Rooktest voor de bankparsers op synthetische voorbeeldbestanden (geen echte gegevens).
Gebruik: python3 tests/test_bankparsers.py   (vereist openpyxl)"""
import shutil, sys, tempfile, subprocess
from decimal import Decimal
from pathlib import Path

KIT = Path(__file__).resolve().parent.parent
TOOLS = KIT / "template" / "Financien" / "_tools"
sys.path.insert(0, str(TOOLS))
import camt053, csv_nl  # noqa: E402

F = KIT / "tests" / "fixtures"
assert csv_nl.herken(F / "ING/ing_voorbeeld.csv") == "ing"
assert csv_nl.herken(F / "Rabobank/rabo_voorbeeld.csv") == "rabobank"
assert csv_nl.herken(F / "ABNAMRO/abn_voorbeeld.TXT") == "abnamro_txt"
assert csv_nl.herken(F / "CAMT/camt_voorbeeld.xml") == "camt053"

ing = csv_nl.ing(F / "ING/ing_voorbeeld.csv")
assert ing[0].bedrag == Decimal("-112.50") and ing[0].incassant_id == "NL00ZZZ000000000001"
assert ing[0].machtigingskenmerk == "KLANT-445566" and ing[1].bedrag == Decimal("1400.00")
rabo = csv_nl.rabobank(F / "Rabobank/rabo_voorbeeld.csv")
assert rabo[1].bedrag == Decimal("-34.95") and rabo[1].machtigingskenmerk == "POLIS-123456"
abn = csv_nl.abnamro_txt(F / "ABNAMRO/abn_voorbeeld.TXT")
assert abn[0].tegenpartij == "Voorbeeld Streaming" and abn[0].incassant_id == "NL00ZZZ000000000003"
assert abn[0].machtigingskenmerk == "ABO-778899" and abn[0].omschrijving == "Abonnement januari"
camt, afschr = camt053.parse(F / "CAMT/camt_voorbeeld.xml")
assert camt[0].bedrag == Decimal("-95.00") and camt[0].incassant_id == "NL00ZZZ000000000004"
assert camt[0].omschrijving == "Premie maart relatienr 990011" and afschr[0]["klopt"] is True

# volledige build in een tijdelijke kopie
with tempfile.TemporaryDirectory() as d:
    fin = Path(d) / "Financien"
    shutil.copytree(TOOLS, fin / "_tools")
    for sub in ("ING", "Rabobank", "ABNAMRO", "CAMT"):
        shutil.copytree(F / sub, fin / sub)
    out = subprocess.run([sys.executable, str(fin / "_tools/bouw_overzicht.py")], capture_output=True, text=True)
    print(out.stdout, out.stderr)
    assert out.returncode == 0 and (fin / "Transacties.xlsx").exists()
    from openpyxl import load_workbook
    wb = load_workbook(fin / "Transacties.xlsx")
    assert wb.sheetnames == ["Transacties", "Terugkerend", "Per maand", "Controle"]
    terug = [r[0] for r in wb["Terugkerend"].iter_rows(min_row=2, values_only=True)]
    assert "Voorbeeld Energie BV" in terug and "Voorbeeld Streaming" in terug, terug
    cats = {r[3]: r[4] for r in wb["Transacties"].iter_rows(min_row=2, values_only=True)}
    assert cats["Eigen rekening"] == "Intern: eigen rekeningen"
print("OK – alle bankparser-tests geslaagd")
