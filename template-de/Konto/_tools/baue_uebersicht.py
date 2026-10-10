#!/usr/bin/env python3
"""Baut Konto/Transaktionen.xlsx aus allen Bankexporten in Konto/<Bank>/.

Aufruf:  .venv/bin/python Konto/_tools/baue_uebersicht.py          (aus dem Repository-Wurzelordner)
Erkennt: camt.053 (XML), Sparkasse CSV-CAMT, DKB-CSV, ING-CSV (Deutschland), Commerzbank-PDF (Giro + Kreditkarte).
Unbekannte Dateien werden gemeldet.

Schritte: einlesen -> entdoppeln -> Salden prüfen (mit Verkettung) -> Lücken melden -> eigene Umbuchungen markieren ->
Anbieter/Kategorie vergeben -> wiederkehrende Zahlungen bestimmen -> Arbeitsmappe schreiben.
Die Mappe wird immer komplett neu gebaut: nie von Hand bearbeiten, Excel vorher schließen.
"""
from __future__ import annotations
import sys
from collections import defaultdict
from decimal import Decimal
from pathlib import Path
from statistics import median

sys.path.insert(0, str(Path(__file__).parent))
import camt053, csv_de, commerzbank_pdf      # noqa: E402
from anbieter import einordnen, INTERN       # noqa: E402

KONTO = Path(__file__).resolve().parent.parent
AUS = KONTO / "Transaktionen.xlsx"


def einlesen(ordner: Path):
    transaktionen, auszuege, pruefungen, unbekannt = [], [], [], []
    for p in sorted(ordner.rglob("*")):
        if not p.is_file() or "_tools" in p.parts or p.name.startswith((".", "~$")) or p.suffix.lower() in (".md", ".xlsx"):
            continue
        art = csv_de.erkenne(p)
        try:
            if art in ("camt053", "commerzbank_pdf"):
                t, a = (camt053 if art == "camt053" else commerzbank_pdf).parse(p)
                auszuege += a
            elif art:
                t = getattr(csv_de, art)(p)
                pruefungen.append(saldo_laufend(p.name, t))
            else:
                unbekannt.append(p.name); continue
        except ValueError as e:
            unbekannt.append(f"{p.name} ({e})"); continue
        transaktionen += t
    return transaktionen, auszuege, pruefungen, unbekannt


def saldo_laufend(name, t):
    """Laufenden Saldo je Zeile prüfen, in Dateireihenfolge (auf- oder absteigend)."""
    mit = [x for x in t if x.saldo_nach is not None]
    if len(mit) < 2:
        return (name, t[0].konto if t else "", "laufender Saldo", None, "keine Salden im Export – nicht prüfbar")
    auf = sum(1 for a, b in zip(mit, mit[1:]) if a.saldo_nach + b.betrag != b.saldo_nach)
    ab = sum(1 for a, b in zip(mit, mit[1:]) if b.saldo_nach + a.betrag != a.saldo_nach)
    f = min(auf, ab)
    return (name, mit[0].konto, "laufender Saldo", f == 0, f"{f} Brüche" if f else "geht auf")


def auszuege_pruefen(auszuege):
    """Anfang + Buchungen = Ende. Geht ein Auszug nicht auf, gegen den Endsaldo des Vorauszugs verketten:
    manche Banken drucken einen falschen Anfangssaldo, obwohl keine Buchung fehlt."""
    zeilen, je_konto = [], defaultdict(list)
    for a in auszuege:
        je_konto[a["konto"]].append(a)
    for konto, liste in je_konto.items():
        liste.sort(key=lambda a: (a["ende"] is None, str(a["id"])))
        vorher = None
        for a in liste:
            if a["stimmt"]:
                zeilen.append((a["quelle"], konto, "Anfang + Buchungen = Ende", True, "geht auf"))
            elif vorher is not None and vorher["ende"] + a["summe"] == a["ende"]:
                zeilen.append((a["quelle"], konto, "über Vorauszug verkettet", True,
                               f"gedruckter Anfangssaldo weicht um {a['anfang'] - vorher['ende']} ab, Buchungen vollständig"))
            else:
                zeilen.append((a["quelle"], konto, "Anfang + Buchungen = Ende", False,
                               f"Differenz {a['anfang'] + a['summe'] - a['ende']}"))
            vorher = a
    return zeilen


def entdoppeln(t):
    gesehen, aus = set(), []
    for x in sorted(t, key=lambda x: (x.konto, x.datum)):
        if x.schluessel() not in gesehen:
            gesehen.add(x.schluessel()); aus.append(x)
    return aus


def luecken(t):
    je = defaultdict(set)
    for x in t:
        je[x.konto].add((x.datum.year, x.datum.month))
    aus = []
    for konto, mon in je.items():
        j, m = min(mon); ende = max(mon)
        while (j, m) < ende:
            if (j, m) not in mon:
                aus.append((konto, f"{j}-{m:02d}"))
            j, m = (j + 1, 1) if m == 12 else (j, m + 1)
    return aus


def rhythmus(tage: float) -> str:
    for grenze, name in ((10, "wöchentlich"), (40, "monatlich"), (100, "vierteljährlich"), (200, "halbjährlich"), (400, "jährlich")):
        if tage <= grenze:
            return name
    return "unregelmäßig"


def wiederkehrend(t):
    gruppen = defaultdict(list)
    for x in t:
        if x.betrag >= 0 or x.kategorie == INTERN:
            continue
        k = (x.glaeubiger_id, x.mandatsreferenz) if x.glaeubiger_id else ("", x.gegenkonto or x.gegenpartei.upper())
        gruppen[k].append(x)
    zeilen = []
    for (gid, ref), xs in gruppen.items():
        if len({(x.datum.year, x.datum.month) for x in xs}) < 2 and not gid:
            continue
        xs.sort(key=lambda x: x.datum)
        abstaende = [(b.datum - a.datum).days for a, b in zip(xs, xs[1:])]
        letzte = xs[-1]
        jahre = max(Decimal((xs[-1].datum - xs[0].datum).days or 365) / 365, Decimal(1))
        zeilen.append([letzte.anbieter, letzte.kategorie, letzte.konto, gid, ref if gid else "", letzte.gegenkonto,
                       len(xs), rhythmus(median(abstaende)) if abstaende else "einmalig (Lastschrift)",
                       xs[0].datum, letzte.datum, float(-letzte.betrag), float(round(sum(-x.betrag for x in xs) / jahre, 2))])
    return sorted(zeilen, key=lambda r: -r[-1])


def schreiben(t, pruef, luek, wied, unbekannt):
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
    wb = Workbook()

    def blatt(titel, kopf, zeilen, tabelle=None, erstes=False):
        ws = wb.active if erstes else wb.create_sheet()
        ws.title = titel
        ws.append(kopf)
        for z in zeilen:
            ws.append(z)
        if tabelle and zeilen:
            tb = Table(displayName=tabelle, ref=f"A1:{get_column_letter(len(kopf))}{len(zeilen) + 1}")
            tb.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(tb)
        for i, k in enumerate(kopf, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(45, len(k) + 4))
        ws.freeze_panes = "A2"

    blatt("Buchungen", ["Datum", "Konto", "Betrag", "Anbieter", "Kategorie", "Gegenpartei", "Gegenkonto", "Verwendungszweck",
                        "Typ", "Gläubiger-ID", "Mandatsreferenz", "Referenz", "Jahr", "Monat", "Quelle"],
          [[x.datum, x.konto, float(x.betrag), x.anbieter, x.kategorie, x.gegenpartei, x.gegenkonto, x.verwendungszweck,
            x.type, x.glaeubiger_id, x.mandatsreferenz, x.referenz, x.datum.year, f"{x.datum:%Y-%m}", x.quelle] for x in t],
          "tblBuchungen", True)
    blatt("Wiederkehrend", ["Anbieter", "Kategorie", "Konto", "Gläubiger-ID", "Mandatsreferenz", "Gegenkonto", "Anzahl",
                            "Rhythmus", "Erste", "Letzte", "Letzter Betrag", "Pro Jahr (Ø)"], wied, "tblWiederkehrend")
    summe, monate = defaultdict(Decimal), sorted({f"{x.datum:%Y-%m}" for x in t})
    for x in t:
        if x.kategorie != INTERN:
            summe[(x.kategorie, f"{x.datum:%Y-%m}")] += x.betrag
    kats = sorted({k for k, _ in summe})
    blatt("Je Monat", ["Kategorie"] + monate, [[k] + [float(summe.get((k, m), 0)) for m in monate] for k in kats])
    blatt("Prüfung", ["Datei", "Konto", "Prüfung", "Stimmt", "Erläuterung"],
          [list(p) for p in pruef] + [[k, "", f"keine Buchungen in {m}", False, "Export fehlt?"] for k, m in luek]
          + [[u, "", "unbekanntes Format", False, "nicht eingelesen"] for u in unbekannt])
    wb.save(AUS)


def main():
    t, auszuege, pruef, unbekannt = einlesen(KONTO)
    vorher = len(t); t = entdoppeln(t)
    pruef = auszuege_pruefen(auszuege) + pruef
    eigene = {x.konto.replace(" ", "") for x in t}
    for x in t:
        x.anbieter, x.kategorie = einordnen(x.gegenpartei, x.verwendungszweck)
        if x.gegenkonto.replace(" ", "") in eigene:
            x.anbieter, x.kategorie = "Eigenes Konto", INTERN
    luek, wied = luecken(t), wiederkehrend(t)
    schreiben(t, pruef, luek, wied, unbekannt)
    print(f"{len(t)} Buchungen ({vorher - len(t)} Dubletten entfernt) aus {len({x.quelle for x in t})} Dateien, "
          f"{len(eigene)} eigene Konten/Karten.")
    for p in pruef:
        if p[3] is False:
            print(f"  ACHTUNG  {p[0]}: {p[2]} – {p[4]}")
    print(f"  {sum(1 for p in pruef if p[2] == 'über Vorauszug verkettet')} Auszüge nur über den Vorauszug verkettet (siehe Blatt Prüfung)")
    for k, m in luek:
        print(f"  LÜCKE    {k}: keine Buchungen in {m}")
    for u in unbekannt:
        print(f"  UNBEKANNT, nicht eingelesen: {u}")
    print(f"{len(wied)} wiederkehrende Zahlungen; {sum(1 for x in t if x.kategorie == 'Sonstiges')} Buchungen in "
          f"„Sonstiges“ (Regeln in anbieter.py ergänzen).")
    print(f"-> {AUS.relative_to(KONTO.parent)}")


if __name__ == "__main__":
    main()
